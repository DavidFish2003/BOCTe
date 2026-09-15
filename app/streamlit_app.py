"""
BOCTe (Breast Oncology Clinical Triage)
Production-ready clinician interface focused exclusively on Single Patient Triage.
"""

from __future__ import annotations

import io
import json
import os
import sys
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, Optional

import docx
from docx.shared import Pt, RGBColor
import numpy as np
import pandas as pd
import plotly.graph_objects as go
import streamlit as st

# Ensure repository root is on sys.path
REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from src.inference import CLINICAL_FEATURE_LABELS, InferenceEngine

# Page Configuration
st.set_page_config(
    page_title="BOCTe (Breast Oncology Clinical Triage)",
    page_icon="🩺",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# Custom Styling for Clinical Executive Theme
st.markdown(
    """
    <style>
    .main {
        background-color: #0b1120;
        color: #f8fafc;
        font-family: 'Segoe UI', -apple-system, BlinkMacSystemFont, Roboto, sans-serif;
    }
    .stApp {
        background: radial-gradient(circle at 15% 15%, #0f172a 0%, #080d1a 100%);
    }
    .clinical-header {
        background: linear-gradient(135deg, #1e293b 0%, #0f172a 100%);
        border-left: 6px solid #38bdf8;
        border-radius: 10px;
        padding: 20px 26px;
        margin-bottom: 24px;
        box-shadow: 0 4px 15px rgba(0,0,0,0.3);
    }
    .clinical-card {
        background: rgba(30, 41, 59, 0.65);
        border: 1px solid rgba(255, 255, 255, 0.08);
        border-radius: 12px;
        padding: 16px 20px;
        margin-bottom: 14px;
        box-shadow: 0 4px 12px rgba(0, 0, 0, 0.2);
    }
    .triage-card-red {
        background: linear-gradient(135deg, rgba(220, 38, 38, 0.22) 0%, rgba(153, 27, 27, 0.35) 100%);
        border: 1.5px solid #ef4444;
        border-radius: 12px;
        padding: 20px 24px;
        margin-bottom: 16px;
    }
    .triage-card-yellow {
        background: linear-gradient(135deg, rgba(217, 119, 6, 0.22) 0%, rgba(146, 64, 14, 0.35) 100%);
        border: 1.5px solid #f59e0b;
        border-radius: 12px;
        padding: 20px 24px;
        margin-bottom: 16px;
    }
    .triage-card-green {
        background: linear-gradient(135deg, rgba(16, 185, 129, 0.22) 0%, rgba(4, 120, 87, 0.35) 100%);
        border: 1.5px solid #10b981;
        border-radius: 12px;
        padding: 20px 24px;
        margin-bottom: 16px;
    }
    .priority-pill {
        display: inline-block;
        padding: 5px 14px;
        font-weight: 800;
        font-size: 0.85rem;
        letter-spacing: 0.06em;
        border-radius: 9999px;
        text-transform: uppercase;
    }
    .pill-red { background-color: #ef4444; color: #ffffff; }
    .pill-yellow { background-color: #f59e0b; color: #000000; }
    .pill-green { background-color: #10b981; color: #ffffff; }
    .metric-value {
        font-size: 1.5rem;
        font-weight: 700;
        color: #f8fafc;
        margin-top: 4px;
    }
    .metric-label {
        font-size: 0.8rem;
        font-weight: 600;
        color: #94a3b8;
        text-transform: uppercase;
        letter-spacing: 0.05em;
    }
    .report-container {
        background: rgba(15, 23, 42, 0.85);
        border: 1px solid rgba(56, 189, 248, 0.25);
        border-radius: 12px;
        padding: 24px 28px;
        margin-top: 18px;
        box-shadow: 0 6px 20px rgba(0, 0, 0, 0.35);
    }
    .report-header {
        border-bottom: 1px solid rgba(255, 255, 255, 0.12);
        padding-bottom: 14px;
        margin-bottom: 18px;
        display: flex;
        justify-content: space-between;
        align-items: center;
    }
    .report-section-title {
        font-size: 0.95rem;
        font-weight: 700;
        color: #38bdf8;
        text-transform: uppercase;
        letter-spacing: 0.06em;
        margin-top: 14px;
        margin-bottom: 8px;
    }
    .report-table {
        width: 100%;
        border-collapse: collapse;
        font-size: 0.9rem;
        margin-bottom: 12px;
    }
    .report-table th {
        background-color: rgba(30, 41, 59, 0.8);
        color: #94a3b8;
        text-align: left;
        padding: 8px 12px;
        border: 1px solid rgba(255, 255, 255, 0.08);
        font-weight: 600;
        font-size: 0.82rem;
        text-transform: uppercase;
    }
    .report-table td {
        padding: 8px 12px;
        border: 1px solid rgba(255, 255, 255, 0.06);
        color: #f1f5f9;
    }
    .status-badge {
        padding: 2px 8px;
        border-radius: 4px;
        font-size: 0.8rem;
        font-weight: 600;
    }
    .badge-positive {
        background: rgba(239, 68, 68, 0.2);
        color: #f87171;
        border: 1px solid rgba(239, 68, 68, 0.4);
    }
    .badge-negative {
        background: rgba(16, 185, 129, 0.15);
        color: #34d399;
        border: 1px solid rgba(16, 185, 129, 0.3);
    }
    </style>
    """,
    unsafe_allow_html=True,
)


@st.cache_resource(show_spinner="Initializing Clinical Inference Engine...")
def get_inference_engine() -> InferenceEngine:
    """Load pre-trained models and preprocessing artifacts."""
    models_dir = Path("models")
    try:
        return InferenceEngine(models_dir=models_dir)
    except Exception:
        from main import train_pipeline

        with st.spinner("Initializing models..."):
            train_pipeline()

        return InferenceEngine(models_dir=models_dir)


def render_clinical_gauge(match_percentage: float, color: str) -> go.Figure:
    """Render animated radial gauge showing Malignancy Concordance Score."""
    fig = go.Figure(
        go.Indicator(
            mode="gauge+number",
            value=match_percentage,
            domain={"x": [0, 1], "y": [0, 1]},
            title={"text": "<b>Malignancy Concordance Score</b>", "font": {"size": 17, "color": "#f8fafc"}},
            number={"suffix": "%", "font": {"size": 40, "color": "#ffffff", "family": "Segoe UI"}},
            gauge={
                "axis": {"range": [0, 100], "tickwidth": 1, "tickcolor": "#64748b", "tickfont": {"color": "#94a3b8"}},
                "bar": {"color": color, "thickness": 0.28},
                "bgcolor": "rgba(15, 23, 42, 0.7)",
                "borderwidth": 1,
                "bordercolor": "rgba(255, 255, 255, 0.1)",
                "steps": [
                    {"range": [0, 40], "color": "rgba(16, 185, 129, 0.16)"},
                    {"range": [40, 70], "color": "rgba(245, 158, 11, 0.16)"},
                    {"range": [70, 100], "color": "rgba(239, 68, 68, 0.22)"},
                ],
                "threshold": {
                    "line": {"color": "#ef4444", "width": 3},
                    "thickness": 0.8,
                    "value": 70.0,
                },
            },
        )
    )
    fig.update_layout(
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        font={"color": "#f8fafc"},
        height=250,
        margin=dict(l=20, r=20, t=35, b=20),
    )
    return fig


def generate_docx_report(
    patient_demographics: dict[str, Any],
    physical_findings: dict[str, bool],
    history_comorbidities: dict[str, bool],
    triage: dict[str, Any],
    concordance_pct: float,
    features_dict: dict[str, Any],
    assessment_metrics: dict[str, Any],
    timestamp: str,
) -> bytes:
    """Generate a formatted Microsoft Word (.docx) document for the Clinical Assessment."""
    doc = docx.Document()

    # Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("CLINICAL ONCOLOGY TRIAGE & CONCORDANCE REPORT")
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(15, 23, 42)

    # Subtitle / Timestamp
    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run(f"Evaluation Timestamp: {timestamp}\nTriage Priority: {triage['code']} PRIORITY ({triage['label']})")
    sub_run.font.size = Pt(10)
    sub_run.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_heading("1. PATIENT DEMOGRAPHICS & PRESENTATION", level=2)
    p1 = doc.add_paragraph()
    p1.add_run(f"• Age: {patient_demographics['age']} years old\n")
    p1.add_run(f"• Biological Sex: {patient_demographics['sex']}\n")
    p1.add_run(f"• Laterality: {patient_demographics['laterality']} Breast\n")
    p1.add_run(f"• Clinical Staging: {patient_demographics['stage']}\n")
    p1.add_run(f"• Initial Presentation Date: {patient_demographics['reg_date']}\n")
    p1.add_run(f"• Evaluation Date: {patient_demographics['diag_date']}\n")
    p1.add_run(f"• Diagnostic Lag: {features_dict.get('Diagnostic_Lag_Days', 0):.0f} days\n")
    p1.add_run(f"• Lifestyle Factors: Smoking: {patient_demographics['smoking']} | Alcohol: {patient_demographics['alcohol']}")

    doc.add_heading("2. PHYSICAL FINDINGS & MEDICAL HISTORY", level=2)
    p2 = doc.add_paragraph()
    p2.add_run(f"• Palpable Breast Mass: {'PRESENT' if physical_findings['lump'] else 'ABSENT'}\n")
    p2.add_run(f"• Nipple / Skin Retraction: {'PRESENT' if physical_findings['retraction'] else 'ABSENT'}\n")
    p2.add_run(f"• Localized Breast Swelling: {'PRESENT' if physical_findings['swelling'] else 'ABSENT'}\n")
    p2.add_run(f"• Breast / Mastalgia Pain: {'PRESENT' if physical_findings['pain'] else 'ABSENT'}\n")
    p2.add_run(f"• Nipple Discharge: {'PRESENT' if physical_findings['discharge'] else 'ABSENT'}\n")
    p2.add_run(f"• Family History of Breast Cancer: {'POSITIVE' if history_comorbidities['fam_breast'] else 'NEGATIVE'}\n")
    p2.add_run(f"• Family History of Other Cancers: {'POSITIVE' if history_comorbidities['fam_other'] else 'NEGATIVE'}\n")
    p2.add_run(f"• Comorbidities: Hypertension ({'YES' if history_comorbidities['htn'] else 'NO'}), Diabetes ({'YES' if history_comorbidities['dm'] else 'NO'}), Peptic Ulcer History ({'YES' if history_comorbidities['pud'] else 'NO'})")

    doc.add_heading("3. MALIGNANCY CONCORDANCE ASSESSMENT", level=2)
    p3 = doc.add_paragraph()
    p3.add_run(f"• Malignancy Concordance Score: {concordance_pct}%\n")
    p3.add_run(f"• Triage Priority Level: {triage['code']} PRIORITY ({triage['label']})\n")
    p3.add_run(f"• Symptom Burden Score: {features_dict.get('Symptom_Severity_Index', 0):.0f} / 5\n")
    p3.add_run(f"• Metabolic Risk Score: {features_dict.get('Metabolic_Risk_Score', 0):.0f} / 3\n")
    p3.add_run(f"• Familial Risk Score: {features_dict.get('Familial_History_Score', 0):.0f} / 2\n")
    p3.add_run(f"• Autoencoder Reconstruction MSE: {assessment_metrics.get('autoencoder_mse', 0):.4f}")

    doc.add_heading("4. CLINICAL IMPRESSION & PROTOCOL ACTION PLAN", level=2)
    p4 = doc.add_paragraph()
    p4.add_run(f"Impression:\n{triage['summary']}\n\n")
    act_run = p4.add_run(f"Recommended Action Plan:\n{triage['action']}")
    act_run.bold = True

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio.getvalue()


def reset_form_state():
    """Reset all input fields in session state to cleared / zero default state."""
    st.session_state["input_age"] = 18
    st.session_state["input_sex"] = "Female"
    st.session_state["input_laterality"] = "Left"
    st.session_state["input_smoking"] = "No"
    st.session_state["input_alcohol"] = "No"
    st.session_state["input_lump"] = False
    st.session_state["input_swelling"] = False
    st.session_state["input_pain"] = False
    st.session_state["input_discharge"] = False
    st.session_state["input_retraction"] = False
    st.session_state["input_fam_breast"] = False
    st.session_state["input_fam_other"] = False
    st.session_state["input_htn"] = False
    st.session_state["input_dm"] = False
    st.session_state["input_pud"] = False
    st.session_state["input_stage"] = "Stage 0"
    st.session_state["input_reg_date"] = date.today()
    st.session_state["input_diag_date"] = date.today()


def main():
    engine = get_inference_engine()

    # Clinical Header
    st.markdown(
        """
        <div class="clinical-header">
            <h1 style="margin:0; font-size:1.85rem; font-weight:800; color:#ffffff; letter-spacing:-0.02em;">
                BOCTe (Breast Oncology Clinical Triage)
            </h1>
        </div>
        """,
        unsafe_allow_html=True,
    )

    # Top Action Bar
    top_col1, top_col2 = st.columns([0.85, 0.15])
    with top_col2:
        if st.button("🔄 Reset Form", use_container_width=True, on_click=reset_form_state):
            st.rerun()

    # Main 2-Column Workstation Layout
    col_in, col_out = st.columns([1.1, 1.0], gap="large")

    with col_in:
        st.markdown("### Fill in the patient's information below")

        with st.container():
            st.markdown("#### Patient Demographics & Lifestyle")
            c1, c2, c3 = st.columns(3)
            with c1:
                age_input = st.slider("Patient Age (Years)", min_value=18, max_value=95, value=18, key="input_age")
            with c2:
                sex_input = st.selectbox("Biological Sex", ["Female", "Male"], index=0, key="input_sex")
            with c3:
                laterality_input = st.selectbox("Laterality", ["Left", "Right", "Bilateral"], index=0, key="input_laterality")

            c4, c5 = st.columns(2)
            with c4:
                smoking_input = st.selectbox("Smoking History", ["No", "Yes"], index=0, key="input_smoking")
            with c5:
                alcohol_input = st.selectbox("Alcohol Consumption", ["No", "Yes"], index=0, key="input_alcohol")

            st.markdown("#### Physical Breast Findings")
            s1, s2, s3 = st.columns(3)
            with s1:
                lump_input = st.checkbox("Palpable Breast Mass", value=False, key="input_lump")
                swelling_input = st.checkbox("Localized Breast Swelling", value=False, key="input_swelling")
            with s2:
                pain_input = st.checkbox("Breast / Mastalgia Pain", value=False, key="input_pain")
                discharge_input = st.checkbox("Nipple Discharge", value=False, key="input_discharge")
            with s3:
                retraction_input = st.checkbox("Nipple / Skin Retraction", value=False, key="input_retraction")

            st.markdown("#### Familial & Medical History")
            f1, f2 = st.columns(2)
            with f1:
                fam_breast_input = st.checkbox("Family History of Breast Cancer", value=False, key="input_fam_breast")
                fam_other_input = st.checkbox("Family History of Other Cancers", value=False, key="input_fam_other")
            with f2:
                htn_input = st.checkbox("Hypertension", value=False, key="input_htn")
                dm_input = st.checkbox("Diabetes Mellitus", value=False, key="input_dm")
                pud_input = st.checkbox("Peptic Ulcer History", value=False, key="input_pud")

            st.markdown("#### Clinical Staging & Presentation Timeline")
            t1, t2, t3 = st.columns(3)
            with t1:
                stage_input = st.selectbox("Clinical Stage", ["Stage 0", "Stage I", "Stage II", "Stage III", "Stage IV"], index=0, key="input_stage")
            with t2:
                reg_date_input = st.date_input("Initial Presentation Date", value=date.today(), key="input_reg_date")
            with t3:
                diag_date_input = st.date_input("Clinical Evaluation Date", value=date.today(), key="input_diag_date")

    # Construct patient dictionary
    patient_dict = {
        "DATE OF REG": str(reg_date_input),
        "Diagnosis Date (main)": str(diag_date_input),
        "Sex": sex_input,
        "Current Age": age_input,
        "Family Hx of Breast Cancer": "Yes" if fam_breast_input else "No",
        "Family Hx of Other Cancers": "Yes" if fam_other_input else "No",
        "Smoking Hx": smoking_input,
        "Alcohol History": alcohol_input,
        "Breast Lump": "Yes" if lump_input else "No",
        "Breast Swelling": "Yes" if swelling_input else "No",
        "Breast/Nipple pain": "Yes" if pain_input else "No",
        "Nipple Discharge": "Yes" if discharge_input else "No",
        "Nipple retraction": "Yes" if retraction_input else "No",
        "Hypertension": "Yes" if htn_input else "No",
        "Diabetes": "Yes" if dm_input else "No",
        "PUD": "Yes" if pud_input else "No",
        "Stage (main)": stage_input,
        "Laterality": laterality_input,
    }

    # Execute Clinical Assessment
    assessment = engine.assess_patient(patient_dict)
    triage = assessment["triage_level"]
    concordance_pct = assessment["concordance_percentage"]
    features_dict = assessment["processed_patient_features"]

    with col_out:
        st.markdown("### 📊 Clinical Triage Assessment")

        # Dynamic Triage Priority Card
        card_class = "triage-card-red" if triage["code"] == "RED" else "triage-card-yellow" if triage["code"] == "YELLOW" else "triage-card-green"
        pill_class = "pill-red" if triage["code"] == "RED" else "pill-yellow" if triage["code"] == "YELLOW" else "pill-green"

        st.markdown(
            f"""
            <div class="{card_class}">
                <div style="display:flex; justify-content:space-between; align-items:center; margin-bottom:10px;">
                    <span class="priority-pill {pill_class}">Priority: {triage['code']}</span>
                    <span style="font-weight:700; font-size:1.1rem; color:#ffffff;">Malignancy Concordance Score: {concordance_pct}%</span>
                </div>
                <div style="font-size:1.2rem; font-weight:800; color:#ffffff; margin-bottom:8px;">{triage['label']}</div>
                <div style="font-size:0.95rem; color:#f1f5f9; margin-bottom:12px; line-height:1.45;">
                    <b>Clinical Rationale:</b> {triage['summary']}
                </div>
                <div style="font-size:0.92rem; color:#e2e8f0; background:rgba(0,0,0,0.25); padding:10px 14px; border-radius:8px; border-left:3px solid {triage['color']};">
                    <b>Recommended Action Plan:</b><br>{triage['action']}
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        # Radial Malignancy Concordance Gauge
        st.plotly_chart(render_clinical_gauge(concordance_pct, triage["color"]), use_container_width=True)

        # Clinical Summary Metrics
        s_c1, s_c2, s_c3, s_c4 = st.columns(4)
        with s_c1:
            st.markdown(
                f"""
                <div class="clinical-card">
                    <div class="metric-label">Symptom Burden</div>
                    <div class="metric-value" style="color:#38bdf8;">{features_dict['Symptom_Severity_Index']:.0f} <span style="font-size:0.85rem; color:#94a3b8;">/ 5</span></div>
                </div>
                """,
                unsafe_allow_html=True,
            )
        with s_c2:
            st.markdown(
                f"""
                <div class="clinical-card">
                    <div class="metric-label">Metabolic Risk</div>
                    <div class="metric-value" style="color:#a78bfa;">{features_dict['Metabolic_Risk_Score']:.0f} <span style="font-size:0.85rem; color:#94a3b8;">/ 3</span></div>
                </div>
                """,
                unsafe_allow_html=True,
            )
        with s_c3:
            st.markdown(
                f"""
                <div class="clinical-card">
                    <div class="metric-label">Familial Risk</div>
                    <div class="metric-value" style="color:#f59e0b;">{features_dict['Familial_History_Score']:.0f} <span style="font-size:0.85rem; color:#94a3b8;">/ 2</span></div>
                </div>
                """,
                unsafe_allow_html=True,
            )
        with s_c4:
            st.markdown(
                f"""
                <div class="clinical-card">
                    <div class="metric-label">Diagnostic Lag</div>
                    <div class="metric-value" style="color:#f472b6;">{features_dict['Diagnostic_Lag_Days']:.0f} <span style="font-size:0.85rem; color:#94a3b8;">d</span></div>
                </div>
                """,
                unsafe_allow_html=True,
            )

    # Clinical Assessment Export Options
    st.markdown("---")
    st.markdown("### 📋 Export Clinical Assessment")

    now_timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    exp_col1, exp_col2 = st.columns(2)

    # Formatted Markdown Clinical Report
    report_markdown = f"""# CLINICAL ONCOLOGY TRIAGE & CONCORDANCE REPORT
Evaluation Timestamp: {now_timestamp}

## 1. PATIENT DEMOGRAPHICS & PRESENTATION
- Age: {age_input} years old
- Biological Sex: {sex_input}
- Laterality: {laterality_input} Breast
- Clinical Staging: {stage_input}
- Initial Presentation Date: {reg_date_input}
- Evaluation Date: {diag_date_input}
- Diagnostic Lag: {features_dict['Diagnostic_Lag_Days']:.0f} days
- Lifestyle Factors: Smoking: {smoking_input} | Alcohol: {alcohol_input}

## 2. PHYSICAL FINDINGS & MEDICAL HISTORY
- Palpable Breast Mass: {"PRESENT" if lump_input else "ABSENT"}
- Nipple / Skin Retraction: {"PRESENT" if retraction_input else "ABSENT"}
- Localized Breast Swelling: {"PRESENT" if swelling_input else "ABSENT"}
- Breast / Mastalgia Pain: {"PRESENT" if pain_input else "ABSENT"}
- Nipple Discharge: {"PRESENT" if discharge_input else "ABSENT"}
- Family History of Breast Cancer: {"POSITIVE" if fam_breast_input else "NEGATIVE"}
- Family History of Other Cancers: {"POSITIVE" if fam_other_input else "NEGATIVE"}
- Comorbidities: Hypertension ({"YES" if htn_input else "NO"}), Diabetes ({"YES" if dm_input else "NO"}), Peptic Ulcer History ({"YES" if pud_input else "NO"})

## 3. MALIGNANCY CONCORDANCE ASSESSMENT
- Malignancy Concordance Score: {concordance_pct}%
- Triage Priority Level: {triage['code']} PRIORITY ({triage['label']})
- Symptom Burden Score: {features_dict['Symptom_Severity_Index']:.0f} / 5
- Metabolic Risk Score: {features_dict['Metabolic_Risk_Score']:.0f} / 3
- Familial Risk Score: {features_dict['Familial_History_Score']:.0f} / 2
- Autoencoder Reconstruction MSE: {assessment['metrics']['autoencoder_mse']:.4f}

## 4. CLINICAL IMPRESSION & PROTOCOL ACTION PLAN
- Impression: {triage['summary']}
- Recommended Action: {triage['action']}
"""

    with exp_col1:
        st.download_button(
            label="📄 Download Formatted Clinical Report (Markdown)",
            data=report_markdown,
            file_name=f"clinical_oncology_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.md",
            mime="text/markdown",
            use_container_width=True,
        )

    # Formatted Microsoft Word (.docx) Clinical Report
    docx_bytes = generate_docx_report(
        patient_demographics={
            "age": age_input,
            "sex": sex_input,
            "laterality": laterality_input,
            "stage": stage_input,
            "reg_date": reg_date_input,
            "diag_date": diag_date_input,
            "smoking": smoking_input,
            "alcohol": alcohol_input,
        },
        physical_findings={
            "lump": lump_input,
            "retraction": retraction_input,
            "swelling": swelling_input,
            "pain": pain_input,
            "discharge": discharge_input,
        },
        history_comorbidities={
            "fam_breast": fam_breast_input,
            "fam_other": fam_other_input,
            "htn": htn_input,
            "dm": dm_input,
            "pud": pud_input,
        },
        triage=triage,
        concordance_pct=concordance_pct,
        features_dict=features_dict,
        assessment_metrics=assessment["metrics"],
        timestamp=now_timestamp,
    )

    with exp_col2:
        st.download_button(
            label="📝 Download Clinical Report (Microsoft Word)",
            data=docx_bytes,
            file_name=f"clinical_oncology_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )


if __name__ == "__main__":
    main()
